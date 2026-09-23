"""Xuất dữ liệu ra CSV / Excel / PDF."""
import csv
import io
from pathlib import Path

from fpdf import FPDF, FontFace
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.ai.service import strip_accents

FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
]

MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM để Excel đọc đúng tiếng Việt


def to_xlsx(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title: str, subtitle: str, sections: list[tuple[str, list[str], list[list], list[float]]]) -> bytes:
    """sections: (tiêu đề, headers, rows, tỉ lệ độ rộng cột)."""
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    font_path = next((p for p in FONT_CANDIDATES if p.exists()), None)
    if font_path:
        pdf.add_font("U", "", str(font_path))
        family, fix = "U", (lambda s: s)
    else:  # không có font Unicode: bỏ dấu tiếng Việt để không lỗi
        family, fix = "Helvetica", strip_accents
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_font(family, size=16)
    pdf.cell(0, 10, fix(title), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font(family, size=10)
    pdf.cell(0, 6, fix(subtitle), new_x="LMARGIN", new_y="NEXT", align="C")
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    for sec_title, headers, rows, ratios in sections:
        pdf.ln(4)
        pdf.set_font(family, size=12)
        pdf.cell(0, 8, fix(sec_title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(family, size=9)
        widths = [usable * r / sum(ratios) for r in ratios]
        heading = FontFace(emphasis="", fill_color=(219, 234, 254))
        with pdf.table(col_widths=widths, text_align="LEFT", line_height=5, headings_style=heading) as table:
            header = table.row()
            for h in headers:
                header.cell(fix(h))
            for r in rows:
                row = table.row()
                for v in r:
                    row.cell(fix(str(v)))
    return bytes(pdf.output())


def vnd(v: int) -> str:
    return f"{v:,.0f}".replace(",", ".")
